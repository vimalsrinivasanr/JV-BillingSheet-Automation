import openpyxl
from collections import Counter

manual_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/manual_output/Output Apr'26.xlsx"
wb = openpyxl.load_workbook(manual_path, data_only=False)
ws = wb.active

pkeys = []
accounts = []
amounts = []
formulas = []

for r in range(2, ws.max_row + 1):
    pk = ws.cell(r, 11).value
    acc = ws.cell(r, 12).value
    amt = ws.cell(r, 10).value
    pkeys.append(pk)
    accounts.append(acc)
    if isinstance(amt, str) and amt.startswith('='):
        formulas.append((r, amt))
    else:
        amounts.append(amt)

print("Unique Posting Keys (data_only=False):", Counter(pkeys))
print("Unique Accounts (data_only=False):", Counter(accounts))
print("Number of formulas found in Amount:", len(formulas))
print("First 5 formulas in Amount:", formulas[:5])
