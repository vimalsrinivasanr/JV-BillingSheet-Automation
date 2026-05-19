import pandas as pd
import sys
import os

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from engine import JVEngine

def test():
    engine = JVEngine()
    normalized_file = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/input file/Input Apr'26_NORMALIZED.xlsx"
    
    df = engine._load_normalized_data(normalized_file)
    
    # Apply filters
    for col in ["workday_id", "classification", "billed_status", "invoice_no"]:
        df[col] = df[col].astype(str).str.strip()
        
    df = df[df["workday_id"].isin(["", "nan", "None"]) == False]
    df = df[
        (df["classification"].str.lower() == "billable") &
        (df["billed_status"].str.lower() == "billed")
    ]
    df = df[~df["invoice_no"].isin(["", "nan", "None"])]
    
    for col in engine.GL_COL_NAMES:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
        
    print("Before aggregation, filtered rows:", len(df))
    
    # Let's aggregate
    # We want to group by invoice_no and employee ID (workday_id, emp_no_ref)
    group_cols = ["invoice_no", "workday_id", "emp_no_ref"]
    agg_dict = {}
    for c in df.columns:
        if c in engine.GL_COL_NAMES:
            agg_dict[c] = "sum"
        elif c not in group_cols:
            agg_dict[c] = "first"
            
    df_agg = df.groupby(group_cols, as_index=False).agg(agg_dict)
    print("After aggregation, total rows:", len(df_agg))
    
    df_agg = df_agg.sort_values(by=["invoice_no", "workday_id"]).reset_index(drop=True)
    rows = engine._build_rows(df_agg)
    
    non_empty_rows = [r for r in rows if r["Reference"] is not None]
    print("Number of non-empty rows in generated JV:", len(non_empty_rows))
    
    # Compare with manual
    manual_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/manual_output/Output Apr'26.xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(manual_path, read_only=True)
    manual_df = pd.read_excel(manual_path, sheet_name=wb.sheetnames[0])
    if "Reference" not in manual_df.columns:
        manual_df = pd.read_excel(manual_path, sheet_name=wb.sheetnames[0], skiprows=2)
    manual_non_empty = manual_df[manual_df["Reference"].notna()]
    print("Manual non-empty rows count:", len(manual_non_empty))
    
if __name__ == "__main__":
    test()
