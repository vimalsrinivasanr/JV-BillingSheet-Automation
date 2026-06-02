import openpyxl
from collections import Counter

manual_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/manual_output/Output.xlsx"

# Load values
wb_val = openpyxl.load_workbook(manual_path, data_only=True)
ws_val = wb_val.active

# Load formulas
wb_form = openpyxl.load_workbook(manual_path, data_only=False)
ws_form = wb_form.active

def get_expected_pk(account, amount):
    acct = str(account or "").strip()
    try:
        amt = float(amount)
    except Exception:
        amt = 0.0
    if len(acct) > 6:
        return "21" if amt > 0 else "31"
    if len(acct) == 6:
        return "40" if amt > 0 else "50"
    return "01" if amt > 0 else "11"

mismatches = []
all_combinations = []
empty_rows = 0

for r in range(4, ws_val.max_row + 1):
    pk_val = ws_val.cell(row=r, column=11).value
    acc_val = ws_val.cell(row=r, column=12).value
    amt_val = ws_val.cell(row=r, column=10).value
    
    pk_formula = ws_form.cell(row=r, column=11).value
    
    # Skip spacer/blank rows
    if pk_val is None and acc_val is None and amt_val is None:
        empty_rows += 1
        continue
        
    expected_pk = get_expected_pk(acc_val, amt_val)
    
    # Store combinations for pattern analysis
    # format: (Account_Length, Amount_Sign, Actual_PK, Expected_PK, PKey_Formula_Present)
    acc_len = len(str(acc_val or "").strip())
    
    try:
        amt_float = float(amt_val) if amt_val is not None else 0.0
    except:
        amt_float = 0.0
        
    amt_sign = "positive" if amt_float > 0 else ("negative" if amt_float < 0 else "zero")
    has_formula = str(pk_formula or "").startswith("=")
    
    all_combinations.append((acc_len, amt_sign, pk_val, expected_pk, has_formula))
    
    # Compare
    pk_val_str = str(pk_val).strip() if pk_val is not None else ""
    expected_pk_str = str(expected_pk).strip()
    
    # Excel could evaluate a number like 40 or 50 as integer, check both string/int formats
    is_match = False
    if pk_val_str == expected_pk_str:
        is_match = True
    elif pk_val_str.replace(".0", "") == expected_pk_str.replace(".0", ""):
        is_match = True
    elif pk_val_str == "1" and expected_pk_str == "01": # "01" is matched to 1 in data_only load
        is_match = True
    
    if not is_match:
        mismatches.append({
            "row": r,
            "actual_pk": pk_val,
            "expected_pk": expected_pk,
            "formula": pk_formula,
            "account": acc_val,
            "amount": amt_val
        })

print(f"Total rows scanned: {ws_val.max_row - 3} (excluding header rows)")
print(f"Empty/spacer rows: {empty_rows}")
print(f"Mismatches found: {len(mismatches)}")

if mismatches:
    print("\n--- Sample Mismatches ---")
    for m in mismatches[:15]:
        print(f"Row {m['row']}: Actual={m['actual_pk']} (Expected={m['expected_pk']}) | Formula={m['formula']} | Account={m['account']} | Amount={m['amount']}")
else:
    print("\nNo mismatches found between python expectation and sheet evaluations.")

# Count unique combinations of Account Length, Amount Sign, Actual PK, and Formula status
print("\n--- Distribution of Patterns (Account Length, Amount Sign -> Actual PK [Expected] | Has Formula?) ---")
pattern_counts = Counter(all_combinations)
for pat, count in pattern_counts.most_common():
    acc_len, amt_sign, pk_val, expected_pk, has_formula = pat
    formula_str = "Formula" if has_formula else "Hardcoded"
    print(f"AccLen={acc_len}, AmtSign={amt_sign} -> Actual PK={pk_val} [Expected: {expected_pk}] | {formula_str} | Count: {count}")
