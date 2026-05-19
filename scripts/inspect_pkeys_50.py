import openpyxl

manual_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/manual_output/Output Apr'26.xlsx"
wb = openpyxl.load_workbook(manual_path)
ws = wb.active

print("Rows with Posting Key = 50:")
count_50 = 0
for r in range(2, ws.max_row + 1):
    pk = ws.cell(r, 11).value
    if pk == 50 or pk == "50":
        row_vals = [ws.cell(r, c).value for c in range(1, 14)]
        print(f"Row {r}: {row_vals}")
        count_50 += 1
        if count_50 >= 10:
            print("Truncated 50 list...")
            break

print("\nRows with Posting Key = None:")
count_none = 0
for r in range(2, ws.max_row + 1):
    pk = ws.cell(r, 11).value
    if pk is None:
        row_vals = [ws.cell(r, c).value for c in range(1, 14)]
        print(f"Row {r}: {row_vals}")
        count_none += 1
        if count_none >= 10:
            print("Truncated None list...")
            break
