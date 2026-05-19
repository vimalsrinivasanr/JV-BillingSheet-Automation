import pandas as pd
import openpyxl
import numpy as np

def load_manual_jv(filepath):
    print(f"Loading manual file: {filepath}...")
    df = pd.read_excel(filepath)
    if "Reference" not in df.columns:
        df = pd.read_excel(filepath, skiprows=2)
    df_clean = df[df["Reference"].notna() & (df["Reference"] != "")].copy()
    df_clean["Reference.1"] = df_clean["Reference.1"].astype(str).str.strip().str.replace(".0", "", regex=False)
    df_clean["Posting Key"] = pd.to_numeric(df_clean["Posting Key"], errors="coerce").fillna(0.0).astype(int)
    df_clean["Account"] = pd.to_numeric(df_clean["Account"], errors="coerce").fillna(0.0).astype(int)
    df_clean["Amount_Evaluated"] = pd.to_numeric(df_clean["Amount"], errors="coerce").fillna(0.0)
    # Normalize Ref Key 3 (20) (Employee ID) to string and clean placeholders
    df_clean["Ref Key 3 (20)"] = df_clean["Ref Key 3 (20)"].astype(str).str.strip().str.replace(".0", "", regex=False)
    df_clean["Ref Key 3 (20)"] = df_clean["Ref Key 3 (20)"].replace(["nan", "None", ""], "0")
    return df_clean

def load_and_evaluate_auto_jv(filepath):
    print(f"Loading and evaluating automated file: {filepath}...")
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=False)
    sheet = wb.active
    
    # Read all rows
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(list(row))
        
    # Find the header row
    header_idx = None
    for idx, row in enumerate(rows):
        if "Reference" in row and "Document Date" in row:
            header_idx = idx
            break
            
    if header_idx is None:
        raise ValueError(f"Could not find header row in {filepath}")
        
    headers = rows[header_idx]
    seen = {}
    clean_headers = []
    for i, h in enumerate(headers):
        h_str = str(h).strip() if h is not None else f"Col{i}"
        if h_str in seen:
            seen[h_str] += 1
            clean_headers.append(f"{h_str}.{seen[h_str]}")
        else:
            seen[h_str] = 0
            clean_headers.append(h_str)
            
    headers = clean_headers
    
    # Process data rows
    data_rows = []
    for r_idx in range(header_idx + 1, len(rows)):
        row = rows[r_idx]
        if len(row) < len(headers):
            row = row + [None] * (len(headers) - len(row))
        else:
            row = row[:len(headers)]
            
        row_dict = dict(zip(headers, row))
        excel_row_num = r_idx + 1
        row_dict["_excel_row"] = excel_row_num
        data_rows.append(row_dict)
        
    memo = {}
    def get_val(row_dict):
        r_num = row_dict["_excel_row"]
        if r_num in memo:
            return memo[r_num]
            
        val = row_dict.get("Amount")
        if val is None:
            res = 0.0
        elif isinstance(val, str) and val.startswith("="):
            clean_formula = val.replace(" ", "").upper()
            if clean_formula.startswith("=-SUM(J") and ":" in clean_formula:
                range_part = clean_formula.replace("=-SUM(J", "").replace(")", "")
                try:
                    start_r, end_r = map(int, range_part.split(":J"))
                    total = 0.0
                    for dr in data_rows:
                        if start_r <= dr["_excel_row"] <= end_r:
                            total += get_val(dr)
                    res = round(-total, 2)
                except Exception as e:
                    print(f"Error evaluating formula {val} on row {r_num}: {e}")
                    res = 0.0
            else:
                res = 0.0
        else:
            try:
                res = float(val)
            except ValueError:
                res = 0.0
        memo[r_num] = res
        return res

    for dr in data_rows:
        dr["Amount_Evaluated"] = get_val(dr)
        
    df = pd.DataFrame(data_rows)
    df_clean = df[df["Reference"].notna() & (df["Reference"] != "")].copy()
    df_clean["Reference.1"] = df_clean["Reference.1"].astype(str).str.strip().str.replace(".0", "", regex=False)
    df_clean["Posting Key"] = pd.to_numeric(df_clean["Posting Key"], errors="coerce").fillna(0.0).astype(int)
    df_clean["Account"] = pd.to_numeric(df_clean["Account"], errors="coerce").fillna(0.0).astype(int)
    # Normalize Ref Key 3 (20) (Employee ID) to string and clean placeholders
    df_clean["Ref Key 3 (20)"] = df_clean["Ref Key 3 (20)"].astype(str).str.strip().str.replace(".0", "", regex=False)
    df_clean["Ref Key 3 (20)"] = df_clean["Ref Key 3 (20)"].replace(["nan", "None", ""], "0")
    return df_clean

def compare():
    manual_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/manual_output/Output Apr'26.xlsx"
    auto_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/input file/SAP_JV_Upload_Run.xlsx"

    manual_df = load_manual_jv(manual_path)
    auto_df = load_and_evaluate_auto_jv(auto_path)

    print(f"\n--- Row Counts ---")
    print(f"Manual non-empty rows: {len(manual_df)}")
    print(f"Automated non-empty rows: {len(auto_df)}")

    print(f"\n--- Totals check ---")
    man_pk40 = manual_df[manual_df["Posting Key"] == 40]
    man_pk50 = manual_df[manual_df["Posting Key"] == 50]
    
    auto_pk40 = auto_df[auto_df["Posting Key"] == 40]
    auto_pk50 = auto_df[auto_df["Posting Key"] == 50]

    print(f"Manual: Posting Key 40 Sum = {man_pk40['Amount_Evaluated'].sum():,.2f} ({len(man_pk40)} rows)")
    print(f"Manual: Posting Key 50 Sum = {man_pk50['Amount_Evaluated'].sum():,.2f} ({len(man_pk50)} rows)")
    
    print(f"Auto:   Posting Key 40 Sum = {auto_pk40['Amount_Evaluated'].sum():,.2f} ({len(auto_pk40)} rows)")
    print(f"Auto:   Posting Key 50 Sum = {auto_pk50['Amount_Evaluated'].sum():,.2f} ({len(auto_pk50)} rows)")

    # Group by invoice and check total credit/debit balances
    print("\n--- Balance check per Invoice ---")
    man_inv_grp = manual_df.groupby("Reference.1")["Amount_Evaluated"].sum()
    auto_inv_grp = auto_df.groupby("Reference.1")["Amount_Evaluated"].sum()
    
    mismatches = []
    for inv in set(man_inv_grp.index).union(auto_inv_grp.index):
        m_val = man_inv_grp.get(inv, 0.0)
        a_val = auto_inv_grp.get(inv, 0.0)
        if abs(m_val - a_val) > 0.05:
            mismatches.append((inv, m_val, a_val, a_val - m_val))
            
    print(f"Invoices with balance mismatches: {len(mismatches)}")
    for m in mismatches[:10]:
        print(f"  Invoice {m[0]} | Manual net sum: {m[1]:,.2f} | Auto net sum: {m[2]:,.2f} | Diff: {m[3]:,.2f}")

    print("\n--- Individual Invoice Net Sums ---")
    man_net_sums = manual_df.groupby("Reference.1")["Amount_Evaluated"].sum().round(2)
    auto_net_sums = auto_df.groupby("Reference.1")["Amount_Evaluated"].sum().round(2)
    
    print(f"Manual invoices not balancing to zero (diff > 0.05): {len(man_net_sums[man_net_sums.abs() > 0.05])}")
    print(f"Auto invoices not balancing to zero (diff > 0.05): {len(auto_net_sums[auto_net_sums.abs() > 0.05])}")
    
    print("\n--- Detailed Row Verification ---")
    mismatched_details = 0
    for inv in manual_df["Reference.1"].unique():
        m_inv = manual_df[manual_df["Reference.1"] == inv].copy()
        a_inv = auto_df[auto_df["Reference.1"] == inv].copy()
        
        m_grp = m_inv.groupby(["Ref Key 3 (20)", "Account"])["Amount_Evaluated"].sum().round(2)
        a_grp = a_inv.groupby(["Ref Key 3 (20)", "Account"])["Amount_Evaluated"].sum().round(2)
        
        diff = m_grp.align(a_grp)
        m_align, a_align = diff[0], diff[1]
        
        m_align = m_align.fillna(0.0)
        a_align = a_align.fillna(0.0)
        
        diff_mask = (m_align - a_align).abs() > 0.05
        if diff_mask.any():
            print(f"  Invoice {inv} has row discrepancies:")
            mismatched_details += 1
            for idx in m_align[diff_mask].index:
                print(f"    Employee {idx[0]} | Account {idx[1]} | Manual: {m_align.loc[idx]:,.2f} | Auto: {a_align.loc[idx]:,.2f}")
                
    print(f"Total invoices with detailed discrepancies: {mismatched_details}")

if __name__ == "__main__":
    compare()
