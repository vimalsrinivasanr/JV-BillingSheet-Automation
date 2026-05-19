import openpyxl

manual_path = "/Users/macbook/Downloads/Library/PROJECTS/Randstad/JV-BillingSheet-Automation/output/manual_output/Output Apr'26.xlsx"
wb = openpyxl.load_workbook(manual_path)
ws = wb.active

print("Offset rows (Account 500003) details:")
for r in range(2, ws.max_row + 1):
    acc = ws.cell(r, 12).value
    if acc == 500003 or acc == "500003":
        pk = ws.cell(r, 11).value
        amt = ws.cell(r, 10).value
        ref = ws.cell(r, 1).value
        print(f"Row {r}: Ref={ref}, PostingKey={pk}, Amount={amt}")
